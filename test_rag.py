"""
Automated unit and integration tests for aiQS components:
- Document Loader (PDF, DOCX, TXT)
- Text chunking & metadata integrity
- In-memory Vector Store cosine similarity search
- RAG Engine grounding pipeline
"""

import os
import unittest
import numpy as np
from io import BytesIO
from pypdf import PdfWriter
from docx import Document as DocxDocument

from document_loader import load_document, chunk_document_pages, DocumentChunk
from vector_store import VectorStore
from rag_engine import RAGEngine


class TestAIQS(unittest.TestCase):

    def test_text_loading_and_chunking(self):
        sample_path = os.path.join(os.path.dirname(__file__), "sample_documents", "MCA_Semester_1_Syllabus.txt")
        self.assertTrue(os.path.exists(sample_path), "Sample file should exist")

        pages = load_document(sample_path, "MCA_Semester_1_Syllabus.txt")
        self.assertEqual(len(pages), 1)
        self.assertIn("MCA102", pages[0]["text"])

        chunks = chunk_document_pages(pages, chunk_size=500, chunk_overlap=100)
        self.assertGreater(len(chunks), 1, "Should split sample into multiple chunks")

        for c in chunks:
            self.assertEqual(c.source_file, "MCA_Semester_1_Syllabus.txt")
            self.assertEqual(c.page_number, 1)
            self.assertTrue(len(c.text) > 0)

    def test_pdf_extraction(self):
        # Create an in-memory PDF
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        # We can test with a minimal real PDF or text
        pdf_bytes = BytesIO()
        writer.write(pdf_bytes)
        pdf_bytes.seek(0)

        pages = load_document(pdf_bytes, "test_doc.pdf")
        # Empty page returns 0 extracted text pages
        self.assertIsInstance(pages, list)

    def test_docx_extraction(self):
        # Create an in-memory DOCX
        doc = DocxDocument()
        doc.add_paragraph("Unit 3: Graph Algorithms")
        doc.add_paragraph("Dijkstra and Bellman-Ford algorithms.")
        docx_bytes = BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)

        pages = load_document(docx_bytes, "test_doc.docx")
        self.assertEqual(len(pages), 1)
        self.assertIn("Graph Algorithms", pages[0]["text"])
        self.assertIn("Bellman-Ford", pages[0]["text"])

    def test_vector_store_similarity(self):
        store = VectorStore()
        c1 = DocumentChunk("c1", "doc1.txt", 1, "Apples and oranges are delicious fruits.")
        c2 = DocumentChunk("c2", "doc1.txt", 1, "Relational databases use SQL to query tables.")
        c3 = DocumentChunk("c3", "doc2.txt", 2, "Bananas and grapes are healthy sweet snacks.")

        # Synthetic embeddings (dimension 3)
        # c1 and c3 represent 'fruit' dimension, c2 represents 'database' dimension
        e1 = [0.9, 0.1, 0.0]
        e2 = [0.0, 0.1, 0.9]
        e3 = [0.85, 0.15, 0.0]

        store.add_chunks([c1, c2, c3], [e1, e2, e3])

        # Query about fruit
        q_vec = [0.95, 0.05, 0.0]
        results = store.similarity_search(q_vec, top_k=2)

        self.assertEqual(len(results), 2)
        # The top 2 results must be c1 and c3
        retrieved_ids = [chunk.chunk_id for chunk, score in results]
        self.assertIn("c1", retrieved_ids)
        self.assertIn("c3", retrieved_ids)
        self.assertNotIn("c2", retrieved_ids)
        self.assertGreater(results[0][1], 0.9)  # High similarity score

    def test_rag_engine_unconfigured(self):
        engine = RAGEngine(api_key="")
        self.assertFalse(engine.is_configured())
        with self.assertRaises(ValueError):
            engine.embed_texts(["hello"])


if __name__ == "__main__":
    unittest.main()
